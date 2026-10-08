"""Generate and reopen real documents in box-doctor's temporary directory."""

import sys
from pathlib import Path
from subprocess import check_output

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

destination = Path(sys.argv[1])
marker = "LumenBox document check"

# Exercise the plotting and image libraries without a display or a writable home.
plt.plot([1, 2, 3], [2, 4, 3])
plt.title(marker)
chart = destination / "chart.png"
plt.savefig(chart)
plt.close()
with Image.open(chart) as image:
    image.verify()

document = Document()
document.add_paragraph(marker + " 文档检查")
document.add_picture(str(chart))
document.save(destination / "check.docx")
assert Document(destination / "check.docx").paragraphs[0].text == marker + " 文档检查"

workbook = Workbook()
workbook.active.append([marker, 42])
workbook.save(destination / "check.xlsx")
reopened = load_workbook(destination / "check.xlsx")
assert reopened.active["A1"].value == marker
assert reopened.active["B1"].value == 42
reopened.close()

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = marker + " 文档检查"
slide.shapes.add_picture(str(chart), Inches(1), Inches(2), width=Inches(5))
presentation.save(destination / "check.pptx")
assert Presentation(destination / "check.pptx").slides[0].shapes.title.text == marker + " 文档检查"

pdf = canvas.Canvas(str(destination / "check.pdf"))
pdf.drawString(72, 750, marker)
pdf.drawImage(str(chart), 72, 350, width=400, height=300)
pdf.save()
assert marker in check_output(["pdftotext", str(destination / "check.pdf"), "-"], text=True)
